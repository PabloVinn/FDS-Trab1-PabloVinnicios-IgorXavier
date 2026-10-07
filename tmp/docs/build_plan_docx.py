from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"E:\igor\Pucrs\FDS-Trab1-PabloVinnicios-IgorXavier")
SOURCE = ROOT / "output" / "docs" / "PLANO_DE_IMPLEMENTACAO_ACMEPOLLING.md"
OUTPUT = ROOT / "output" / "docs" / "Plano_de_Implementacao_ACMEPolling.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=110, bottom=110, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_font(run, name="Aptos", size=None, bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, "Consolas", 9.5)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, bold=True)
        else:
            run = paragraph.add_run(part)
            set_font(run)


def add_checkbox(doc, text, checked):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.55)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.keep_together = True
    box = "☑" if checked else "☐"
    run = p.add_run(f"{box} ")
    set_font(run, "Arial", 10.5)
    add_inline(p, text)
    return p


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.footer_distance = Cm(0.55)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title_style.font.size = Pt(25)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_p_pr = title_style.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size, before, after in (("Heading 1", 17, 14, 6), ("Heading 2", 13, 10, 4), ("Heading 3", 11, 8, 3)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    i = 0
    first_paragraph = True
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_pr = p._p.get_or_add_pPr()
            existing_border = p_pr.find(qn("w:pBdr"))
            if existing_border is not None:
                p_pr.remove(existing_border)
            add_inline(p, line[2:])
            i += 1
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=2)
            i += 1
            continue

        if line.startswith("```"):
            lang = line[3:].strip()
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.right_indent = Cm(0.25)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.keep_together = False
            for idx, code_line in enumerate(code_lines):
                run = p.add_run(code_line)
                set_font(run, "Consolas", 8.5)
                if idx < len(code_lines) - 1:
                    run.add_break()
            continue

        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1]):
            headers = [c.strip() for c in line.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(headers))
            table.autofit = False
            set_table_borders(table)
            widths = [Cm(2.2), Cm(4.0), Cm(10.0)] if len(headers) == 3 else [Cm(16.2 / len(headers))] * len(headers)
            for col, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
                cell.width = widths[col]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_shading(cell, "1F4E78")
                set_cell_margins(cell)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(header)
                set_font(r, bold=True, color="FFFFFF")
            for ridx, row_data in enumerate(rows):
                cells = table.add_row().cells
                for col, (cell, value) in enumerate(zip(cells, row_data)):
                    cell.width = widths[col]
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    set_cell_margins(cell)
                    if ridx % 2:
                        set_cell_shading(cell, "EAF2F8")
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    add_inline(p, value)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        checkbox = re.match(r"^- \[([ xX])\] (.*)$", line)
        if checkbox:
            add_checkbox(doc, checkbox.group(2), checkbox.group(1).lower() == "x")
            i += 1
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
            i += 1
            continue

        numbered = re.match(r"^(\d+)\. (.*)$", line)
        if numbered:
            if numbered.group(1) == "1":
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(0)
                spacer.paragraph_format.line_spacing = Pt(2)
            p = doc.add_paragraph(style="List Number")
            add_inline(p, numbered.group(2))
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, line)
        if first_paragraph:
            p.paragraph_format.space_after = Pt(10)
            first_paragraph = False
        i += 1

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ACMEPolling   Plano de implementação")
    set_font(r, size=8, color="666666")

    doc.core_properties.title = "Plano de implementação do ACMEPolling"
    doc.core_properties.subject = "Checklist, ordem de implementação e tutorial de execução"
    doc.core_properties.author = "Equipe ACMEPolling"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
